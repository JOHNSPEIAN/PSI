"""Парсинг билетов из Word."""
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ticket_generator.tickets import (TicketsFileError, load_tickets,
                                      parse_ticket_lines)
from tests.helpers import TempDirTestCase, ticket_block, write_tickets


class ParseTicketsFromWordTest(TempDirTestCase):
    def test_parses_all_tickets_with_three_questions(self):
        path = write_tickets(self.tmp / "tickets.docx",
                             ticket_block(1) + [""] + ticket_block(2) + [""] + ticket_block(3))
        tickets = load_tickets(path)
        self.assertEqual(list(tickets), [1, 2, 3])
        self.assertEqual(tickets[2].questions, ("Вопрос 2.1", "Вопрос 2.2", "Вопрос 2.3"))

    def test_question_number_prefix_is_stripped(self):
        tickets = parse_ticket_lines(["Билет 7", "1. Что такое ООП?", "2) Наследование", "3.Инкапсуляция"])
        self.assertEqual(tickets[7].questions, ("Что такое ООП?", "Наследование", "Инкапсуляция"))

    def test_header_variants(self):
        lines = ["Билет №4", "1. a", "2. b", "3. c", "  билет 5: ", "1. d", "2. e", "3. f"]
        self.assertEqual(list(parse_ticket_lines(lines)), [4, 5])

    def test_word_auto_numbered_list_is_supported(self):
        """Вопросы, оформленные автонумерацией Word (номер не входит в текст)."""
        doc = Document()
        doc.add_paragraph("Билет 1")
        for text in ("Первый", "Второй", "Третий"):
            p = doc.add_paragraph(text)
            p_pr = p._p.get_or_add_pPr()
            num_pr = OxmlElement("w:numPr")
            for tag, val in (("w:ilvl", "0"), ("w:numId", "1")):
                el = OxmlElement(tag)
                el.set(qn("w:val"), val)
                num_pr.append(el)
            p_pr.append(num_pr)
        path = self.tmp / "auto.docx"
        doc.save(path)
        self.assertEqual(load_tickets(path)[1].questions, ("Первый", "Второй", "Третий"))

    def test_ticket_with_less_than_three_questions_is_skipped_and_logged(self):
        path = write_tickets(self.tmp / "tickets.docx",
                             ticket_block(1) + ticket_block(2, count=2) + ticket_block(3))
        with self.assertLogs("ticket_generator.tickets", level="ERROR") as logs:
            tickets = load_tickets(path)
        self.assertEqual(list(tickets), [1, 3])
        self.assertTrue(any("Билет 2 пропущен" in m for m in logs.output))

    def test_unrecognized_header_is_skipped_and_logged(self):
        lines = ticket_block(1) + ["Билет X", "1. a", "2. b", "3. c"] + ticket_block(3)
        with self.assertLogs("ticket_generator.tickets", level="ERROR") as logs:
            tickets = parse_ticket_lines(lines)
        self.assertEqual(list(tickets), [1, 3])
        self.assertTrue(any("Билет X" in m for m in logs.output))
        # вопросы нераспознанного билета не попали в соседние
        self.assertEqual(tickets[1].questions, ("Вопрос 1.1", "Вопрос 1.2", "Вопрос 1.3"))

    def test_duplicate_ticket_number_keeps_first(self):
        lines = ticket_block(1, prefix="A") + ticket_block(1, prefix="B")
        with self.assertLogs("ticket_generator.tickets", level="ERROR"):
            tickets = parse_ticket_lines(lines)
        self.assertEqual(tickets[1].questions[0], "A 1.1")

    def test_extra_questions_are_truncated_to_three(self):
        tickets = parse_ticket_lines(ticket_block(1, count=5))
        self.assertEqual(len(tickets[1].questions), 3)

    def test_questions_before_first_header_are_ignored(self):
        tickets = parse_ticket_lines(["1. сирота"] + ticket_block(1))
        self.assertEqual(list(tickets), [1])

    def test_missing_file_raises_clear_error(self):
        with self.assertRaises(TicketsFileError) as ctx:
            load_tickets(self.tmp / "nope.docx")
        self.assertIn("не найден", str(ctx.exception))

    def test_file_without_valid_tickets_raises(self):
        path = write_tickets(self.tmp / "tickets.docx", ticket_block(1, count=1))
        with self.assertLogs("ticket_generator.tickets", level="ERROR"):
            with self.assertRaises(TicketsFileError):
                load_tickets(path)

    def test_corrupted_file_raises_clear_error(self):
        path = self.tmp / "tickets.docx"
        path.write_bytes(b"not a docx")
        with self.assertRaises(TicketsFileError):
            load_tickets(path)


if __name__ == "__main__":
    unittest.main()
