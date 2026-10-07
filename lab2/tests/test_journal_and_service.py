"""Журнал results.xlsx, назначение билета и идемпотентность."""
import random
import unittest
from datetime import datetime, timedelta
from unittest import mock

from openpyxl import load_workbook

from ticket_generator.journal import (HEADERS, JournalLockedError, JournalRow,
                                      ResultsJournal)
from ticket_generator.service import SelectionState, TicketService
from ticket_generator.students import Student, load_students
from ticket_generator.tickets import parse_ticket_lines
from tests.helpers import GROUPS, TempDirTestCase, ticket_block, write_students


def make_tickets(n=5):
    lines = []
    for i in range(1, n + 1):
        lines += ticket_block(i)
    return parse_ticket_lines(lines)


class FakeClock:
    def __init__(self):
        self.now = datetime(2026, 10, 7, 9, 0, 0)

    def __call__(self):
        self.now += timedelta(seconds=1)
        return self.now


def sheet_values(path):
    wb = load_workbook(path)
    try:
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


IVANOV = Student("IA-2301", "Иванов", "Иван")
PETROVA = Student("IA-2301", "Петрова", "Анна")


class JournalTest(TempDirTestCase):
    def setUp(self):
        super().setUp()
        self.journal = ResultsJournal(self.tmp / "results.xlsx")

    def test_created_with_header_if_missing(self):
        self.assertTrue(self.journal.ensure_exists())
        self.assertEqual(sheet_values(self.journal.path), [list(HEADERS)])
        self.assertFalse(self.journal.ensure_exists())  # второй раз не пересоздаётся

    def test_append_keeps_existing_rows_untouched(self):
        t = datetime(2026, 1, 1, 10, 0, 0)
        self.journal.append(JournalRow("G", "А", "Б", 1, t, False))
        before = sheet_values(self.journal.path)
        self.journal.append(JournalRow("G", "В", "Г", 2, t, False))
        after = sheet_values(self.journal.path)
        self.assertEqual(after[:len(before)], before)
        self.assertEqual(after[-1], ["G", "В", "Г", 2, t, "нет"])

    def test_find_first_ticket_uses_earliest_row(self):
        t = datetime(2026, 1, 1)
        self.journal.append(JournalRow("G", "Иванов", "Иван", 4, t, False))
        self.journal.append(JournalRow("G", "Иванов", "Иван", 9, t, True))  # подделка «руками»
        self.assertEqual(self.journal.find_first_ticket("G", "Иванов", "Иван"), 4)
        self.assertEqual(self.journal.find_first_ticket(" g ", "ИВАНОВ", "иван"), 4)
        self.assertIsNone(self.journal.find_first_ticket("Другая", "Иванов", "Иван"))

    def test_locked_file_on_write_raises_locked_error(self):
        self.journal.ensure_exists()
        before = sheet_values(self.journal.path)
        with mock.patch("ticket_generator.journal.open", side_effect=PermissionError, create=True):
            with self.assertRaises(JournalLockedError) as ctx:
                self.journal.append(JournalRow("G", "А", "Б", 1, datetime.now(), False))
        self.assertIn("открыт в Excel", str(ctx.exception))
        self.assertEqual(sheet_values(self.journal.path), before)

    def test_locked_file_on_replace_raises_locked_error_and_leaves_no_temp(self):
        self.journal.ensure_exists()
        with mock.patch("ticket_generator.journal.os.replace", side_effect=PermissionError):
            with self.assertRaises(JournalLockedError):
                self.journal.append(JournalRow("G", "А", "Б", 1, datetime.now(), False))
        self.assertEqual([p.name for p in self.tmp.iterdir()], ["results.xlsx"])


class TicketServiceTest(TempDirTestCase):
    def setUp(self):
        super().setUp()
        self.tickets = make_tickets(5)
        self.journal = ResultsJournal(self.tmp / "results.xlsx")
        self.journal.ensure_exists()
        self.service = TicketService(self.tickets, self.journal, rng=random.Random(42), clock=FakeClock())

    def test_new_student_gets_random_ticket_from_file(self):
        seen = set()
        for i in range(30):
            a = self.service.generate(Student("G", f"Студент{i}", "Тест"))
            self.assertIn(a.ticket_no, self.tickets)
            self.assertEqual(a.questions, self.tickets[a.ticket_no].questions)
            self.assertEqual(len(a.questions), 3)
            self.assertFalse(a.repeat)
            seen.add(a.ticket_no)
        self.assertGreater(len(seen), 1, "билеты должны выбираться случайно")

    def test_random_choice_only_among_valid_tickets(self):
        tickets = parse_ticket_lines(ticket_block(1) + ticket_block(2, count=1) + ticket_block(3))
        service = TicketService(tickets, self.journal, rng=random.Random(0))
        numbers = {service.assign(Student("G", f"S{i}", "N")).ticket_no for i in range(50)}
        self.assertEqual(numbers, {1, 3})

    def test_repeat_selection_returns_same_ticket(self):
        """Идемпотентность: повторный выбор того же студента → тот же билет."""
        first = self.service.generate(IVANOV)
        for _ in range(5):
            again = self.service.generate(IVANOV)
            self.assertEqual(again.ticket_no, first.ticket_no)
            self.assertEqual(again.questions, first.questions)
            self.assertTrue(again.repeat)

    def test_repeat_appends_new_row_with_yes(self):
        first = self.service.generate(IVANOV)
        rows_before = sheet_values(self.journal.path)
        self.service.generate(IVANOV)
        rows_after = sheet_values(self.journal.path)
        self.assertEqual(len(rows_after), len(rows_before) + 1)
        self.assertEqual(rows_after[:len(rows_before)], rows_before)  # старые строки не тронуты
        self.assertEqual(rows_after[-1][:4], ["IA-2301", "Иванов", "Иван", first.ticket_no])
        self.assertEqual(rows_after[-1][5], "да")
        self.assertEqual(rows_before[-1][5], "нет")

    def test_repeat_survives_restart(self):
        first = self.service.generate(IVANOV)
        restarted = TicketService(self.tickets, ResultsJournal(self.journal.path), rng=random.Random(999))
        self.assertEqual(restarted.generate(IVANOV).ticket_no, first.ticket_no)

    def test_same_name_in_other_group_is_different_student(self):
        self.service.generate(IVANOV)
        other = self.service.assign(Student("IA-2302", "Иванов", "Иван"))
        self.assertFalse(other.repeat)

    def test_row_is_written_immediately_after_generation(self):
        self.service.generate(PETROVA)
        self.assertEqual(sheet_values(self.journal.path)[-1][1:3], ["Петрова", "Анна"])

    def test_full_history_in_chronological_order(self):
        order = [IVANOV, PETROVA, IVANOV, PETROVA, IVANOV]
        for s in order:
            self.service.generate(s)
        rows = self.journal.rows()
        self.assertEqual([(r.surname, r.repeat) for r in rows],
                         [("Иванов", False), ("Петрова", False), ("Иванов", True),
                          ("Петрова", True), ("Иванов", True)])
        stamps = [r.timestamp for r in rows]
        self.assertEqual(stamps, sorted(stamps))
        self.assertTrue(all(isinstance(t, datetime) for t in stamps))

    def test_retry_after_lock_saves_same_ticket(self):
        assignment = self.service.assign(IVANOV)
        with mock.patch("ticket_generator.journal.os.replace", side_effect=PermissionError):
            with self.assertRaises(JournalLockedError):
                self.service.save(assignment)
        self.assertEqual(self.journal.rows(), [])
        self.service.save(assignment)  # «Повторить» после закрытия Excel
        self.assertEqual(self.journal.find_first_ticket("IA-2301", "Иванов", "Иван"), assignment.ticket_no)

    def test_repeat_of_ticket_removed_from_file(self):
        self.journal.append(JournalRow("IA-2301", "Иванов", "Иван", 99, datetime.now(), False))
        a = self.service.assign(IVANOV)
        self.assertEqual((a.ticket_no, a.questions, a.repeat), (99, (), True))


class SelectionStateTest(TempDirTestCase):
    def setUp(self):
        super().setUp()
        self.state = SelectionState(load_students(write_students(self.tmp / "students.xlsx", GROUPS)))

    def test_initially_button_disabled_and_groups_listed(self):
        self.assertEqual(self.state.group_names, ["IA-2301", "IA-2302", "IA-2303"])
        self.assertFalse(self.state.can_generate)

    def test_group_selection_fills_students_of_that_sheet(self):
        self.assertEqual(self.state.select_group("IA-2302"), ["Морозов Алексей", "Волкова Елена"])
        self.assertFalse(self.state.can_generate)  # только группа — кнопка неактивна

    def test_button_enabled_only_when_both_selected(self):
        self.state.select_group("IA-2301")
        self.state.select_student("Петрова Анна")
        self.assertTrue(self.state.can_generate)
        self.state.select_group("IA-2302")  # смена группы сбрасывает студента
        self.assertFalse(self.state.can_generate)

    def test_empty_group_gives_empty_list_and_disabled_button(self):
        self.assertEqual(self.state.select_group("IA-2303"), [])
        self.state.select_student("Иванов Иван")
        self.assertFalse(self.state.can_generate)

    def test_student_without_group_or_unknown_does_not_enable(self):
        self.state.select_student("Иванов Иван")
        self.assertFalse(self.state.can_generate)
        self.state.select_group("IA-2301")
        self.state.select_student("Волкова Елена")  # из другой группы
        self.assertFalse(self.state.can_generate)

    def test_reset_student_disables_button(self):
        self.state.select_group("IA-2301")
        self.state.select_student("Иванов Иван")
        self.state.reset_student()
        self.assertFalse(self.state.can_generate)
        self.assertEqual(self.state.group, "IA-2301")


if __name__ == "__main__":
    unittest.main()
