"""Вспомогательные функции для создания тестовых файлов."""
from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook


def write_students(path: Path, groups: dict[str, list[tuple]]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for group, rows in groups.items():
        ws = wb.create_sheet(group)
        ws.append(["Фамилия", "Имя"])
        for row in rows:
            ws.append(list(row))
    wb.save(path)
    return path


def write_tickets(path: Path, paragraphs: list[str]) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)
    return path


def ticket_block(number: int, count: int = 3, prefix: str = "Вопрос") -> list[str]:
    return [f"Билет {number}"] + [f"{i}. {prefix} {number}.{i}" for i in range(1, count + 1)]


GROUPS = {
    "IA-2301": [("Иванов", "Иван"), ("Петрова", "Анна"), ("Сидоров", "Павел")],
    "IA-2302": [("Морозов", "Алексей"), ("Волкова", "Елена")],
    "IA-2303": [],
}


class TempDirTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="tickets-test-"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)
