"""Журнал результатов results.xlsx.

Колонки: Группа | Фамилия | Имя | Номер билета | Дата и время | Повтор (да/нет)

Данные только дописываются в конец: существующие строки не изменяются,
поэтому по журналу восстанавливается полная хронология всех генераций.
"""
from __future__ import annotations

import logging
import os
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

log = logging.getLogger(__name__)

HEADERS = ("Группа", "Фамилия", "Имя", "Номер билета", "Дата и время", "Повтор (да/нет)")
SHEET_TITLE = "Журнал"
DATETIME_FORMAT = "DD.MM.YYYY HH:MM:SS"
REPEAT_YES = "да"
REPEAT_NO = "нет"


class JournalError(Exception):
    """Журнал не читается (повреждён, не xlsx и т.п.)."""


class JournalLockedError(JournalError):
    """Журнал недоступен для записи — обычно открыт в Excel."""


@dataclass(frozen=True)
class JournalRow:
    group: str
    surname: str
    name: str
    ticket_no: int
    timestamp: datetime
    repeat: bool

    def as_cells(self) -> list:
        return [self.group, self.surname, self.name, self.ticket_no, self.timestamp,
                REPEAT_YES if self.repeat else REPEAT_NO]


def _norm(value) -> str:
    return "" if value is None else str(value).strip().casefold()


class ResultsJournal:
    def __init__(self, path: str | Path):
        self.path = Path(path)

    # ---------- создание ----------
    def ensure_exists(self) -> bool:
        """Создаёт пустой журнал с заголовком, если файла нет. Возвращает True, если создан."""
        if self.path.exists():
            return False
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET_TITLE
        ws.append(list(HEADERS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCDEF", (14, 18, 14, 14, 21, 16)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        self._save(wb)
        log.info("Создан журнал %s", self.path)
        return True

    # ---------- чтение ----------
    def _open(self, read_only: bool):
        try:
            return load_workbook(self.path, read_only=read_only)
        except PermissionError as exc:
            raise JournalLockedError(self._locked_message()) from exc
        except Exception as exc:
            raise JournalError(f"Не удалось прочитать {self.path.name}: {exc}") from exc

    def rows(self) -> list[JournalRow]:
        """Все записи журнала в порядке следования (хронологическом)."""
        if not self.path.exists():
            return []
        wb = self._open(read_only=True)
        result: list[JournalRow] = []
        try:
            ws = wb.worksheets[0]
            for row in ws.iter_rows(min_row=2, max_col=6, values_only=True):
                row = tuple(row) + (None,) * (6 - len(row))
                if all(v is None or str(v).strip() == "" for v in row):
                    continue
                try:
                    ticket_no = int(row[3])
                except (TypeError, ValueError):
                    log.warning("Журнал: строка с некорректным номером билета %r пропущена", row)
                    continue
                result.append(JournalRow(
                    group=str(row[0] or "").strip(),
                    surname=str(row[1] or "").strip(),
                    name=str(row[2] or "").strip(),
                    ticket_no=ticket_no,
                    timestamp=row[4],
                    repeat=_norm(row[5]) == REPEAT_YES,
                ))
        finally:
            wb.close()
        return result

    def find_first_ticket(self, group: str, surname: str, name: str) -> int | None:
        """Номер билета из самой первой записи по студенту или None."""
        key = (_norm(group), _norm(surname), _norm(name))
        for row in self.rows():
            if (_norm(row.group), _norm(row.surname), _norm(row.name)) == key:
                return row.ticket_no
        return None

    # ---------- запись ----------
    def _locked_message(self) -> str:
        return (f"Файл {self.path.name} недоступен для записи — возможно, он открыт в Excel. "
                f"Закройте файл и повторите попытку.")

    def _check_writable(self) -> None:
        """Проверяет, что файл можно открыть на запись (в Windows Excel держит блокировку)."""
        if not self.path.exists():
            return
        try:
            with open(self.path, "r+b"):
                pass
        except PermissionError as exc:
            raise JournalLockedError(self._locked_message()) from exc

    def _save(self, wb: Workbook) -> None:
        """Атомарное сохранение: во временный файл рядом, затем замена."""
        fd, tmp_name = tempfile.mkstemp(prefix=".~results-", suffix=".xlsx", dir=self.path.parent)
        os.close(fd)
        try:
            wb.save(tmp_name)
            os.replace(tmp_name, self.path)
        except PermissionError as exc:
            raise JournalLockedError(self._locked_message()) from exc
        finally:
            if os.path.exists(tmp_name):
                os.remove(tmp_name)

    def append(self, row: JournalRow) -> None:
        """Дописывает строку в конец журнала и сразу сохраняет файл."""
        self.ensure_exists()
        self._check_writable()
        wb = self._open(read_only=False)
        ws = wb.worksheets[0]
        ws.append(row.as_cells())
        ws.cell(row=ws.max_row, column=5).number_format = DATETIME_FORMAT
        self._save(wb)
        log.info("Журнал: %s %s %s — билет %d, повтор=%s",
                 row.group, row.surname, row.name, row.ticket_no, row.repeat)
