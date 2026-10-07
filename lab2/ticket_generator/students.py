"""Чтение списка студентов из students.xlsx.

Формат: один лист на группу, имя листа = имя группы.
Строка 1 — заголовок, данные со строки 2: A — Фамилия, B — Имя.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

log = logging.getLogger(__name__)


@dataclass(frozen=True)
class Student:
    group: str
    surname: str
    name: str

    @property
    def display_name(self) -> str:
        """Формат для выпадающего списка: «Фамилия Имя»."""
        return f"{self.surname} {self.name}"


class StudentsFileError(Exception):
    """students.xlsx отсутствует или не читается."""


def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def load_students(path: str | Path) -> dict[str, list[Student]]:
    """Возвращает {имя группы: [Student, ...]} в порядке листов книги.

    Пустая группа (лист без студентов) остаётся в словаре с пустым списком.
    """
    path = Path(path)
    if not path.is_file():
        raise StudentsFileError(f"Файл со списком студентов не найден: {path}")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # повреждённый файл, не xlsx и т.п.
        raise StudentsFileError(f"Не удалось прочитать {path.name}: {exc}") from exc

    groups: dict[str, list[Student]] = {}
    try:
        for ws in wb.worksheets:
            group = ws.title.strip()
            students: list[Student] = []
            seen: set[tuple[str, str]] = set()
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
                surname = _cell_text(row[0] if len(row) > 0 else None)
                name = _cell_text(row[1] if len(row) > 1 else None)
                if not surname and not name:
                    continue  # пустая строка
                if not surname or not name:
                    log.warning("Группа %s, строка %d: не заполнены фамилия или имя — пропущено",
                                group, row_idx)
                    continue
                key = (surname.casefold(), name.casefold())
                if key in seen:
                    log.warning("Группа %s, строка %d: дубликат студента %s %s — пропущено",
                                group, row_idx, surname, name)
                    continue
                seen.add(key)
                students.append(Student(group, surname, name))
            groups[group] = students
    finally:
        wb.close()
    return groups


def find_student(groups: dict[str, list[Student]], group: str, display_name: str) -> Student | None:
    """Поиск студента по группе и строке «Фамилия Имя» из выпадающего списка."""
    for student in groups.get(group, []):
        if student.display_name == display_name:
            return student
    return None
