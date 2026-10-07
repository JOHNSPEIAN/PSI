"""Создаёт демонстрационные data/students.xlsx и data/tickets.docx.

Запуск: python tools/make_sample_data.py [ПАПКА]
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

GROUPS = {
    "IA-2301": [("Иванов", "Иван"), ("Петрова", "Анна"), ("Сидоров", "Павел"),
                ("Кузнецова", "Мария"), ("Попов", "Дмитрий")],
    "IA-2302": [("Морозов", "Алексей"), ("Волкова", "Елена"), ("Лебедев", "Никита"),
                ("Соколова", "Ольга")],
    "IA-2303": [],  # пустая группа — для проверки
}

QUESTIONS = [
    "Типы данных в Python: изменяемые и неизменяемые.",
    "Списки и кортежи: отличия и случаи применения.",
    "Словари и множества: устройство и основные операции.",
    "Функции: позиционные и именованные аргументы, *args и **kwargs.",
    "Области видимости переменных, правило LEGB.",
    "Генераторы списков и выражения-генераторы.",
    "Обработка исключений: try/except/else/finally.",
    "Работа с файлами, контекстный менеджер with.",
    "Модули и пакеты, оператор import.",
    "Классы и объекты, метод __init__.",
    "Наследование и полиморфизм.",
    "Инкапсуляция: соглашения об именах в Python.",
    "Декораторы функций.",
    "Итераторы и протокол итерации.",
    "Генераторы и оператор yield.",
    "Модуль unittest: структура тестов.",
    "Виртуальные окружения и pip.",
    "Работа с Git: commit, branch, merge.",
    "Чтение и запись Excel-файлов с помощью openpyxl.",
    "Работа с документами Word через python-docx.",
    "Основы Tkinter: окна и виджеты.",
    "Менеджеры геометрии Tkinter: pack, grid, place.",
    "Обработка событий в Tkinter, метод bind.",
    "Регулярные выражения, модуль re.",
    "Модуль datetime: даты и время.",
    "Логирование, модуль logging.",
    "Аннотации типов и модуль typing.",
    "dataclasses: назначение и использование.",
    "Модуль random: генерация случайных чисел.",
    "Принципы SOLID на примерах Python.",
]


def make_students(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for group, students in GROUPS.items():
        ws = wb.create_sheet(group)
        ws.append(["Фамилия", "Имя"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for surname, name in students:
            ws.append([surname, name])
        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 14
    wb.save(path)


def make_tickets(path: Path) -> None:
    doc = Document()
    for n in range(len(QUESTIONS) // 3):
        doc.add_paragraph(f"Билет {n + 1}")
        for i, q in enumerate(QUESTIONS[n * 3:n * 3 + 3], start=1):
            doc.add_paragraph(f"{i}. {q}")
        doc.add_paragraph("")
    doc.save(path)


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "data"
    target.mkdir(parents=True, exist_ok=True)
    make_students(target / "students.xlsx")
    make_tickets(target / "tickets.docx")
    print(f"Созданы {target / 'students.xlsx'} и {target / 'tickets.docx'}")
