from contextlib import redirect_stdout
from datetime import datetime
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from main import HEADERS, ExitRequested, JournalError, StudentRecord, append_record, read_name, run


def keys(text):
    return iter(text).__next__


class JournalTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.path = Path(self.folder.name) / "journal.xlsx"
        self.record = StudentRecord("Иванов", "Иван", 7, datetime(2026, 9, 2, 12, 30, 15))

    def rows(self):
        book = load_workbook(self.path)
        try:
            return list(book.worksheets[0].values)
        finally:
            book.close()

    def test_new_journal_has_header_number_and_excel_datetime(self):
        append_record(self.path, self.record)
        self.assertEqual(self.rows(), [HEADERS, ("Иванов", "Иван", 7, self.record.created_at)])

    def test_three_then_two_preserves_first_three(self):
        for _ in range(3):
            append_record(self.path, self.record)
        previous = self.rows()
        for _ in range(2):
            append_record(self.path, self.record)
        self.assertEqual(len(self.rows()), 6)
        self.assertEqual(self.rows()[:4], previous)

    def test_names_are_text_not_formulas(self):
        append_record(self.path, StudentRecord("=1+1", "=SUM(1,2)", 1, self.record.created_at))
        book = load_workbook(self.path)
        try:
            self.assertEqual(book.active["A2"].data_type, "s")
            self.assertEqual(book.active["B2"].value, "=SUM(1,2)")
        finally:
            book.close()

    def test_blocked_replace_preserves_original_bytes(self):
        append_record(self.path, self.record)
        original = self.path.read_bytes()
        with patch("main.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaisesRegex(JournalError, "Закройте файл в Excel"):
                append_record(self.path, self.record)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.path.parent.glob(".journal-*")), [])

    def test_failed_write_preserves_original_bytes(self):
        append_record(self.path, self.record)
        original = self.path.read_bytes()
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(JournalError):
                append_record(self.path, self.record)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual(list(self.path.parent.glob(".journal-*")), [])

    def test_excel_owner_file_is_reported(self):
        self.path.with_name("~$journal.xlsx").touch()
        with self.assertRaisesRegex(JournalError, "Закройте файл в Excel"):
            append_record(self.path, self.record)

    def test_corrupt_file_is_not_overwritten(self):
        self.path.write_bytes(b"not an xlsx")
        with self.assertRaises(JournalError):
            append_record(self.path, self.record)
        self.assertEqual(self.path.read_bytes(), b"not an xlsx")

    def test_wrong_headers_are_not_overwritten(self):
        append_record(self.path, self.record)
        book = load_workbook(self.path)
        book.active["A1"] = "Other data"
        book.save(self.path)
        book.close()
        original = self.path.read_bytes()
        with self.assertRaisesRegex(JournalError, "Неверная шапка"):
            append_record(self.path, self.record)
        self.assertEqual(self.path.read_bytes(), original)

    def test_retry_keeps_same_record_and_does_not_duplicate(self):
        original_append = append_record
        attempts = []

        def fail_once(path, record):
            attempts.append(record)
            if len(attempts) == 1:
                raise JournalError("Закройте файл в Excel")
            original_append(path, record)

        output = io.StringIO()
        with patch("main.append_record", side_effect=fail_once), redirect_stdout(output):
            run(self.path, keys("  Иванов  \r Иван \r\r\x1b"))
        self.assertEqual(len(self.rows()), 2)
        self.assertEqual(self.rows()[1][:2], ("Иванов", "Иван"))
        self.assertIs(attempts[0], attempts[1])
        self.assertIn("Запись пока НЕ сохранена", output.getvalue())

    def test_escape_at_first_name_does_not_save_partial_student(self):
        with redirect_stdout(io.StringIO()):
            run(self.path, keys("Иванов\rИв\x1b"))
        self.assertFalse(self.path.exists())

    def test_escape_at_retry_keeps_previous_rows(self):
        append_record(self.path, self.record)
        with patch("main.append_record", side_effect=JournalError("locked")), redirect_stdout(io.StringIO()):
            run(self.path, keys("Петров\rПётр\r\x1b"))
        self.assertEqual(len(self.rows()), 2)

    def test_ticket_range_including_both_endpoints(self):
        with patch("main.secrets.randbelow", side_effect=[0, 19]) as random, redirect_stdout(io.StringIO()):
            run(self.path, keys("A\rB\rC\rD\r\x1b"))
        self.assertEqual([row[2] for row in self.rows()[1:]], [1, 20])
        self.assertEqual([call.args for call in random.call_args_list], [(20,), (20,)])


class InputTests(unittest.TestCase):
    def test_empty_input_and_spaces_repeat_prompt(self):
        output = io.StringIO()
        with redirect_stdout(output):
            result = read_name("Last name: ", keys("\r   \r  Петров  \r"))
        self.assertEqual(result, "Петров")
        self.assertEqual(output.getvalue().count("Поле не может быть пустым"), 2)

    def test_backspace(self):
        with redirect_stdout(io.StringIO()):
            self.assertEqual(read_name("", keys("\bИвах\bн\r")), "Иван")

    def test_escape_immediately_and_mid_input(self):
        for text in ("\x1b", "Ива\x1b"):
            with self.subTest(text=text), redirect_stdout(io.StringIO()), self.assertRaises(ExitRequested):
                read_name("", keys(text))


if __name__ == "__main__":
    unittest.main()
