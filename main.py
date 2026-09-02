"""Консольный генератор экзаменационных билетов. Python 3.10+."""

from __future__ import annotations

import argparse
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime
import os
from pathlib import Path
import secrets
import sys
import tempfile
from typing import Callable, Iterator
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException


HEADERS = ("Last name", "First name", "Номер билета", "Дата и время")


class ExitRequested(Exception):
    """Пользователь нажал ESC или завершил ввод."""


class JournalError(Exception):
    """Журнал нельзя прочитать или сохранить без потери данных."""


@dataclass(frozen=True)
class StudentRecord:
    last_name: str
    first_name: str
    ticket: int
    created_at: datetime


@contextmanager
def keyboard() -> Iterator[Callable[[], str]]:
    """Считывать клавиши немедленно; вернуть настройки терминала при выходе."""
    if not sys.stdin.isatty():
        raise RuntimeError("Запустите программу в терминале с поддержкой клавиатуры.")
    if os.name == "nt":
        import msvcrt

        def read_windows_key() -> str:
            char = msvcrt.getwch()
            if char in ("\x00", "\xe0"):
                msvcrt.getwch()  # Стрелки и функциональные клавиши: второй код.
                return ""
            return char

        yield read_windows_key
    else:
        import select
        import termios
        import tty

        fd = sys.stdin.fileno()
        previous = termios.tcgetattr(fd)
        # Байтовое чтение не оставляет скрытый буфер TextIOWrapper перед select.
        def read_posix_key() -> str:
            first = os.read(fd, 1)
            if not first:
                raise ExitRequested
            if first == b"\x1b":
                # Отличать одиночный ESC от ANSI-последовательностей стрелок.
                if not select.select([fd], [], [], 0.04)[0]:
                    return "\x1b"
                second = os.read(fd, 1)
                if second in (b"[", b"O"):
                    while select.select([fd], [], [], 0.04)[0]:
                        byte = os.read(fd, 1)
                        if not byte or 0x40 <= byte[0] <= 0x7E:
                            break
                    return ""
                raise ExitRequested
            data = bytearray(first)
            length = (1 if first[0] < 0x80 else 2 if first[0] < 0xE0
                      else 3 if first[0] < 0xF0 else 4)
            for _ in range(length - 1):
                byte = os.read(fd, 1)
                if not byte:
                    raise ExitRequested
                data.extend(byte)
            return data.decode("utf-8", errors="replace")

        try:
            tty.setcbreak(fd)
            yield read_posix_key
        finally:
            termios.tcsetattr(fd, termios.TCSADRAIN, previous)


def read_name(prompt: str, read_key: Callable[[], str]) -> str:
    while True:
        print(prompt, end="", flush=True)
        chars: list[str] = []
        while True:
            char = read_key()
            if char in ("\x1b", "\x03", "\x04"):
                raise ExitRequested
            if char in ("\r", "\n"):
                print()
                value = "".join(chars).strip()
                if value:
                    return value
                print("Поле не может быть пустым. Повторите ввод.")
                break
            if char in ("\b", "\x7f"):
                if chars:
                    chars.pop()
                    print("\b \b", end="", flush=True)
            elif char and char.isprintable():
                chars.append(char)
                print(char, end="", flush=True)


def append_record(path: Path, record: StudentRecord) -> None:
    """Добавить строку; существующий файл заменяется только готовым XLSX."""
    workbook = None
    temporary: Path | None = None
    try:
        # Office оставляет такой файл, пока книга открыта. Дополнительная
        # проверка полезна на системах, где Excel не блокирует os.replace.
        if path.with_name("~$" + path.name).exists():
            raise PermissionError("Журнал открыт в Excel")
        if path.exists():
            workbook = load_workbook(path)
            sheet = workbook.worksheets[0]
            if tuple(sheet.cell(1, col).value for col in range(1, 5)) != HEADERS:
                raise JournalError("Неверная шапка журнала. Ожидаются: " + ", ".join(HEADERS))
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Журнал"
            sheet.append(HEADERS)
            sheet.freeze_panes = "A2"
            for column, width in (("A", 24), ("B", 24), ("C", 18), ("D", 23)):
                sheet.column_dimensions[column].width = width

        row = sheet.max_row + 1
        # Явный строковый тип: имя, начинающееся с '=', не станет формулой.
        for column, name in ((1, record.last_name), (2, record.first_name)):
            cell = sheet.cell(row, column, name)
            cell.data_type = "s"
        sheet.cell(row, 3, record.ticket)
        sheet.cell(row, 4, record.created_at).number_format = "yyyy-mm-dd hh:mm:ss"
        with tempfile.NamedTemporaryFile(
            dir=path.parent, prefix=".journal-", suffix=".xlsx", delete=False
        ) as stream:
            temporary = Path(stream.name)
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        workbook.close()
        workbook = None
        os.replace(temporary, path)
        temporary = None
    except PermissionError as error:
        raise JournalError(
            f"Нет доступа к {path.name}. Закройте файл в Excel "
            "и проверьте права на запись в папку."
        ) from error
    except (BadZipFile, InvalidFileException, KeyError, ValueError) as error:
        raise JournalError(
            f"Не удалось прочитать {path.name}: файл повреждён или имеет неверный формат. "
            "Восстановите корректный журнал; существующий файл не заменён."
        ) from error
    except OSError as error:
        raise JournalError(
            f"Не удалось сохранить {path.name}: {error}. "
            "Проверьте папку, свободное место и доступ к файлу."
        ) from error
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass


def run(path: Path, read_key: Callable[[], str]) -> None:
    print("Для выхода нажмите ESC.")
    try:
        while True:
            last_name = read_name("Last name: ", read_key)
            first_name = read_name("First name: ", read_key)
            record = StudentRecord(last_name, first_name, secrets.randbelow(20) + 1, datetime.now())
            print(f"Билет № {record.ticket}.")
            while True:
                try:
                    append_record(path, record)
                    print("Запись сохранена.\n")
                    break
                except JournalError as error:
                    print(f"{error}\nЗапись пока НЕ сохранена.")
                    print("Устраните причину и нажмите Enter для повтора; ESC — выход без этой записи.")
                    while True:
                        char = read_key()
                        if char in ("\x1b", "\x03", "\x04"):
                            raise ExitRequested
                        if char in ("\r", "\n"):
                            break
    except (ExitRequested, KeyboardInterrupt):
        print("\nРабота завершена.")


def main() -> int:
    parser = argparse.ArgumentParser(description="Генератор экзаменационных билетов 1–20")
    parser.add_argument("--journal", type=Path, default=Path(__file__).resolve().with_name("journal.xlsx"),
                        help="путь журнала (по умолчанию journal.xlsx рядом с main.py)")
    args = parser.parse_args()
    try:
        with keyboard() as read_key:
            run(args.journal.expanduser().resolve(), read_key)
    except (RuntimeError, KeyboardInterrupt) as error:
        print(f"\n{error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
